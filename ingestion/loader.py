from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

import fitz


@dataclass
class PageContent:
    page: int
    text: str
    tables: list[list[list[str | None]]] = field(default_factory=list)


@dataclass
class LoadedDocument:
    path: Path
    pages: list[PageContent]


def _load_pdf(path: Path) -> LoadedDocument:
    pages: list[PageContent] = []
    doc = fitz.open(path)
    try:
        table_pages: dict[int, list] = {}
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, start=1):
                    try:
                        table_pages[i] = page.extract_tables() or []
                    except Exception:
                        table_pages[i] = []
        except Exception:
            table_pages = {}

        for i, page in enumerate(doc, start=1):
            pages.append(PageContent(page=i, text=page.get_text("text") or "", tables=table_pages.get(i, [])))
    finally:
        doc.close()
    return LoadedDocument(path=path, pages=pages)


def _load_markdown_or_text(path: Path) -> LoadedDocument:
    return LoadedDocument(path=path, pages=[PageContent(page=1, text=path.read_text(encoding="utf-8"))])


def _load_csv(path: Path) -> LoadedDocument:
    rows: list[list[str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows.extend(list(csv.reader(f)))
    text = "\n".join(" | ".join(row) for row in rows)
    return LoadedDocument(path=path, pages=[PageContent(page=1, text=text, tables=[rows])])


def _load_docx(path: Path) -> LoadedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to ingest DOCX files") from exc
    doc = Document(path)
    text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    tables = []
    for t in doc.tables:
        tables.append([[cell.text for cell in row.cells] for row in t.rows])
    return LoadedDocument(path=path, pages=[PageContent(page=1, text=text, tables=tables)])


def load_document(path: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix in {".md", ".txt"}:
        return _load_markdown_or_text(path)
    if suffix == ".csv":
        return _load_csv(path)
    if suffix == ".docx":
        return _load_docx(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")
